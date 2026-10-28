"""Excel export for the admin: every request, customer and loan in one file."""
import io
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from core import loans as L
from core.store import REQUEST_KINDS

IST = timezone(timedelta(hours=5, minutes=30))
STATUS_LABEL = {"new": "New", "contacted": "Contacted", "in_progress": "In progress",
                "done": "Done", "closed": "Closed"}
RUPEES = '"₹"#,##0'
HEAD_FILL = PatternFill("solid", fgColor="2E3D9E")


def _when(value):
    """ISO timestamp (UTC) -> India time, as a real Excel date."""
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return str(value)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(IST).replace(tzinfo=None)


def _num(value):
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _pct(x):
    try:
        return f"{float(x):g}%"
    except (TypeError, ValueError):
        return ""


def _summary(kind, d):
    """One line on what the customer wants, like the admin page shows."""
    if kind == "balance_transfer" and d.get("loan"):
        loan = d["loan"]
        s = f"{L.LOAN_TYPES.get(loan.get('type'), loan.get('type'))} {L.inr(loan.get('outstanding') or 0)} at {_pct(loan.get('rate'))}"
        offer = d.get("offer")
        if offer:
            s += f" → {offer.get('bank')} from {_pct(offer.get('best', {}).get('rate'))}"
        return s
    if kind == "fresh_loan":
        s = L.LOAN_TYPES.get(d.get("loan_type"), "Loan")
        if d.get("amount_wanted"):
            s += f", wants {d['amount_wanted']}"
        return s
    if kind == "payoff_help" and d.get("plan"):
        p = d["plan"]
        return f"Extra {L.inr(p.get('extra') or 0)}/month, {p.get('months_saved')} months sooner"
    return ""


def _sheet(wb, title, headers, rows, money_cols=(), date_cols=(), first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append([h for h, _ in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            # Anything a customer typed stays plain text, never a formula.
            if isinstance(cell.value, str) and cell.value[:1] in ("=", "+", "-", "@"):
                cell.data_type = "s"
    for i, (_, width) in enumerate(headers, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        if i in money_cols:
            for c in ws[letter][1:]:
                c.number_format = RUPEES
        if i in date_cols:
            for c in ws[letter][1:]:
                c.number_format = "dd mmm yyyy, hh:mm"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def build_workbook(store):
    profiles = {p["user_id"]: p for p in store.admin_all_profiles()}
    loans = store.admin_all_loans()
    requests_ = store.admin_all_requests()

    wb = Workbook()

    # Requests: the sheet you work from
    rows = []
    for r in requests_:
        p = profiles.get(r.get("user_id"), {})
        d = r.get("details") or {}
        offer = d.get("offer") or {}
        rows.append([
            _when(r.get("created_at")),
            STATUS_LABEL.get(r.get("status"), r.get("status")),
            REQUEST_KINDS.get(r.get("kind"), r.get("kind")),
            p.get("name") or "",
            p.get("phone") or d.get("phone") or "",
            p.get("email") or "",
            p.get("employer") or "",
            p.get("city") or "",
            _num(p.get("salary") or d.get("salary")),
            _summary(r.get("kind"), d),
            _num((offer.get("best") or {}).get("net_saving")),
            _num((offer.get("cautious") or {}).get("net_saving")),
            d.get("note") or "",
            r.get("admin_note") or "",
        ])
    _sheet(wb, "Requests", [
        ("Received (IST)", 20), ("Status", 12), ("Request", 22), ("Name", 22), ("Phone", 16),
        ("Email", 28), ("Employer", 20), ("City", 14), ("Salary / month", 15), ("What they want", 48),
        ("Saving (best)", 14), ("Saving (at least)", 16), ("Customer's note", 36), ("Your notes", 36),
    ], rows, money_cols=(9, 11, 12), date_cols=(1,), first=True)

    # Customers: one row each, with their loan totals
    by_user = {}
    for l in loans:
        by_user.setdefault(l.get("user_id"), []).append(l)
    asked = {}
    for r in requests_:
        asked[r.get("user_id")] = asked.get(r.get("user_id"), 0) + 1
    rows = []
    for uid, p in profiles.items():
        mine = by_user.get(uid, [])
        rows.append([
            p.get("name") or "", p.get("phone") or "", p.get("email") or "", p.get("employer") or "",
            p.get("city") or "", _num(p.get("salary")), len(mine),
            sum(_num(l.get("outstanding")) or 0 for l in mine) or None,
            sum(_num(l.get("emi")) or 0 for l in mine) or None,
            asked.get(uid, 0), _when(p.get("created_at")),
        ])
    _sheet(wb, "Customers", [
        ("Name", 22), ("Phone", 16), ("Email", 28), ("Employer", 20), ("City", 14),
        ("Salary / month", 15), ("Loans", 8), ("Total left to pay", 18), ("Total EMI / month", 18),
        ("Requests", 10), ("Joined (IST)", 20),
    ], rows, money_cols=(6, 8, 9), date_cols=(11,))

    # Loans: every loan customers added
    rows = []
    for l in loans:
        p = profiles.get(l.get("user_id"), {})
        rows.append([
            p.get("name") or "", p.get("email") or "", L.LOAN_TYPES.get(l.get("type"), l.get("type") or ""),
            l.get("lender") or "", _num(l.get("outstanding")), _num(l.get("rate")), _num(l.get("emi")),
            l.get("months_left"), _when(l.get("created_at")),
        ])
    _sheet(wb, "Loans", [
        ("Customer", 22), ("Email", 28), ("Loan type", 18), ("Lender", 20), ("Left to pay", 15),
        ("Rate % / year", 13), ("EMI / month", 14), ("Months left", 12), ("Added (IST)", 20),
    ], rows, money_cols=(5, 7), date_cols=(9,))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
